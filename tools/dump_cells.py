import openpyxl, sys, io
P = "reference/analiz pro.2.xlsx"
wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)

out = io.StringIO()
def w(s=""):
    out.write(str(s) + "\n")

for ws in wbf.worksheets:
    wsv = wbv[ws.title]
    w("\n" + "="*100)
    w(f"SHEET: {ws.title} | dims={ws.dimensions}")
    w("="*100)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None: continue
            v = wsv[c.coordinate].value
            if isinstance(c.value, str) and c.value.startswith("="):
                w(f"{c.coordinate}\tFORMULA\t{c.value}\tCACHED={v!r}")
            else:
                w(f"{c.coordinate}\tVALUE\t{c.value!r}")
    # merged + comments
    if ws.merged_cells.ranges:
        w("MERGED: " + ", ".join(str(r) for r in ws.merged_cells.ranges))
    for row in ws.iter_rows():
        for c in row:
            if c.comment:
                w(f"COMMENT {c.coordinate}: {c.comment.text!r} by {c.comment.author}")

txt = out.getvalue()
open("audit/cells_dump.txt","w").write(txt)
print("chars:", len(txt), "lines:", txt.count("\n"))
