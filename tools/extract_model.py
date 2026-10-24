import openpyxl, json
P = "reference/analiz pro.2.xlsx"
wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)

for name in ["کابینت زمینی","کابینت هوایی"]:
    ws, wsv = wbf[name], wbv[name]
    print("="*80); print(name)
    # find rows where column B holds the literal cabinet-type label
    for r in range(1, ws.max_row+1):
        b = ws.cell(r,2).value
        if isinstance(b,str) and b.strip().startswith("کابینت"):
            print(f"  header row {r}: B={b!r}  A={ws.cell(r,1).value!r}")

print("\n"+"="*80)
# drawer sheet detail
ws = wbf["کشو"]
for row in ws.iter_rows(min_row=1, max_row=8, max_col=12):
    print([ (c.coordinate, c.value) for c in row if c.value is not None ])

print("\n"+"="*80)
# check for cells with number formats / rounding
for name in wbf.sheetnames:
    ws = wbf[name]
    fmts = set()
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.number_format != "General":
                fmts.add((c.number_format, c.coordinate))
    print(name, "custom formats:", sorted(fmts)[:20], "count", len(fmts))
