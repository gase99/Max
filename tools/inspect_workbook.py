import openpyxl, json, zipfile, sys
P = "reference/analiz pro.2.xlsx"

# raw zip contents
with zipfile.ZipFile(P) as z:
    print("=== ZIP ENTRIES ===")
    for n in z.namelist():
        print(f"  {n}  ({z.getinfo(n).file_size} B)")

wb = openpyxl.load_workbook(P, data_only=False)
print("\n=== SHEETS ===")
for ws in wb.worksheets:
    print(f"  {ws.title!r} state={ws.sheet_state} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")

print("\n=== DEFINED NAMES ===")
try:
    for name in wb.defined_names:
        dn = wb.defined_names[name]
        print(f"  {name} -> {dn.attr_text}")
except Exception as e:
    print("  err", e)

print("\n=== DATA VALIDATIONS ===")
for ws in wb.worksheets:
    dv = list(ws.data_validations.dataValidation)
    if dv:
        print(f"  [{ws.title}]")
        for d in dv:
            print(f"    type={d.type} formula1={d.formula1} sqref={d.sqref}")

print("\n=== VBA/MACROS ===")
print("  macro-enabled:", P.endswith('.xlsm'))
