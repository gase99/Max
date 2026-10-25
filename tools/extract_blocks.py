import openpyxl, json
P = "reference/analiz pro.2.xlsx"
wbf = openpyxl.load_workbook(P, data_only=False)
wbv = openpyxl.load_workbook(P, data_only=True)

def v(ws,r,c):
    x = ws.cell(r,c).value
    return x

def num(x):
    if x is None: return None
    if isinstance(x,(int,float)): return x
    return x

out = {"ground":[], "wall":[], "drawer":None}

# ---- GROUND: 11 blocks, header rows ----
ws, wsv = wbf["کابینت زمینی"], wbv["کابینت زمینی"]
headers = [1,11,20,30,38,46,54,62,70,78,86]
for i,h in enumerate(headers):
    d = h+1   # data row
    l = h+2   # 'نوار' labels
    f = h+3   # flags
    c_ = h+4  # counts
    blk = {
      "index": i+1,
      "headerRow": h,
      "title": v(ws,h,1),
      "B_width": v(ws,d,2),
      # pairs: (colA=عرض, colB=طول)
      "parts": {
        "base":   {"w": v(ws,d,3), "l": v(ws,d,4), "flag_w": v(ws,f,3), "flag_l": v(ws,f,4), "count": v(ws,c_,4)},
        "body":   {"w": v(ws,d,5), "l": v(ws,d,6), "flag_w": v(ws,f,5), "flag_l": v(ws,f,6), "count": v(ws,c_,6)},
        "shelf":  {"w": v(ws,d,7), "l": v(ws,d,8), "flag_w": v(ws,f,7), "flag_l": v(ws,f,8), "count": v(ws,c_,8)},
        "rail":   {"w": v(ws,d,9), "l": v(ws,d,10), "flag_w": v(ws,f,9), "flag_l": v(ws,f,10), "count": v(ws,c_,10)},
      },
      "band3":  {"w": v(ws,d,11), "l": v(ws,d,12), "count": v(ws,c_,12)},
      "door":   {"w": v(ws,d,13), "h_in": v(ws,d,14), "flag_w": v(ws,f,13), "flag_h": v(ws,f,14), "count": v(ws,c_,14)},
      # excel cached outputs
      "excel0": True,
      "excel": {k: v(wsv,d,cc) for k,cc in [("B",2),("baseW",3),("baseL",4),("bodyW",5),("bodyL",6),
                 ("shelfW",7),("shelfL",8),("railW",9),("railL",10),("band3W",11),("band3L",12),("doorW",13),("doorH",14)]},
    }
    blk["excel"]["doorCount"] = v(wsv,c_,14)
    out["ground"].append(blk)

# ---- WALL: 20 blocks ----
ws, wsv = wbf["کابینت هوایی"], wbv["کابینت هوایی"]
headers = [2,9,16,23,30,37,44,51,58,65,72,79,86,93,100,107,114,121,128,135]
for i,h in enumerate(headers):
    d, f, c_ = h+1, h+3, h+4
    blk = {
      "index": i+1, "headerRow": h,
      "A_row": v(ws,h,1),
      "B_width": v(ws,d,2),
      "parts": {
        "body":  {"w": v(ws,d,3),  "l": v(ws,d,4),  "flag_w": v(ws,f,3),  "flag_l": v(ws,f,4),  "count": v(ws,c_,4)},
        "base":  {"w": v(ws,d,5),  "l": v(ws,d,6),  "flag_w": v(ws,f,5),  "flag_l": v(ws,f,6),  "count": v(ws,c_,6)},
        "top":   {"w": v(ws,d,7),  "l": v(ws,d,8),  "flag_w": v(ws,f,7),  "flag_l": v(ws,f,8),  "count": v(ws,c_,8)},
        "shelf": {"w": v(ws,d,9),  "l": v(ws,d,10), "flag_w": v(ws,f,9),  "flag_l": v(ws,f,10), "count": v(ws,c_,10)},
        "rail":  {"w": v(ws,d,11), "l": v(ws,d,12), "flag_w": v(ws,f,11), "flag_l": v(ws,f,12), "count": v(ws,c_,12)},
      },
      "band3": {"w": v(ws,d,13), "l": v(ws,d,14), "count_w": v(ws,c_,13), "count_l": v(ws,c_,14)},
      "door":  {"w": v(ws,d,15), "h_in": v(ws,d,16), "count": v(ws,c_,16)},
      "excel": {k: v(wsv,d,cc) for k,cc in [("B",2),("bodyW",3),("bodyL",4),("baseW",5),("baseL",6),
                 ("topW",7),("topL",8),("shelfW",9),("shelfL",10),("railW",11),("railL",12),
                 ("band3W",13),("band3L",14),("doorW",15),("doorH",16)]},
    }
    blk["excel"]["doorCount"] = v(wsv,c_,16)
    out["wall"].append(blk)

# ---- DRAWER ----
ws, wsv = wbf["کشو"], wbv["کشو"]
out["drawer"] = {
 "count": v(ws,4,2), "depth": v(ws,4,3), "width": v(ws,4,4), "height": v(ws,4,5),
 "excel": {k: v(wsv,r,c) for k,(r,c) in [("bodyL",(4,8)),("bodyW",(5,8)),("bodyCount",(6,8)),
           ("frontL",(4,10)),("frontW",(5,10)),("frontCount",(6,10)),
           ("backL",(4,12)),("backW",(5,12)),("backCount",(6,12))]},
 "formulas": {"bodyL":v(ws,4,8),"bodyW":v(ws,5,8),"bodyCount":v(ws,6,8),
              "frontL":v(ws,4,10),"frontW":v(ws,5,10),"frontCount":v(ws,6,10),
              "backL":v(ws,4,12),"backW":v(ws,5,12),"backCount":v(ws,6,12)},
}

json.dump(out, open("audit/workbook_model.json","w"), ensure_ascii=False, indent=1)
print("ground blocks:", len(out["ground"]), "wall blocks:", len(out["wall"]))
print("\nGROUND block 1:", json.dumps(out["ground"][0], ensure_ascii=False, indent=1))
print("\nWALL block 1:", json.dumps(out["wall"][0], ensure_ascii=False, indent=1))
